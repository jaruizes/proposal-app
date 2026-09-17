from io import BytesIO

import pytest
from docx import Document
from pypdf import PdfWriter

from agent_platform.application.document_parsers import (
    DOCX_MEDIA_TYPE,
    DocumentParseError,
    DocumentParserRegistry,
)


def test_text_parser_decodes_utf8_bom() -> None:
    parsed = DocumentParserRegistry().parse(
        b"\xef\xbb\xbfhello\nworld",
        filename="notes.txt",
        declared_media_type="text/plain",
    )
    assert parsed.content == "hello\nworld"
    assert parsed.media_type == "text/plain"
    assert parsed.metadata["parser"] == "text"


def test_docx_parser_extracts_paragraphs_and_tables() -> None:
    document = Document()
    document.add_paragraph("Architecture reference")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Platform"
    table.cell(0, 1).text = "OpenShift"
    buffer = BytesIO()
    document.save(buffer)

    parsed = DocumentParserRegistry().parse(
        buffer.getvalue(),
        filename="architecture.docx",
        declared_media_type=DOCX_MEDIA_TYPE,
    )

    assert "Architecture reference" in parsed.content
    assert "Platform | OpenShift" in parsed.content
    assert parsed.media_type == DOCX_MEDIA_TYPE
    assert parsed.metadata["parser"] == "python-docx"


def test_pdf_without_extractable_text_reports_ocr_requirement() -> None:
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    buffer = BytesIO()
    writer.write(buffer)

    with pytest.raises(DocumentParseError, match="require OCR"):
        DocumentParserRegistry().parse(
            buffer.getvalue(),
            filename="scan.pdf",
            declared_media_type="application/pdf",
        )


def test_legacy_doc_is_rejected_explicitly() -> None:
    with pytest.raises(DocumentParseError, match="Legacy .doc"):
        DocumentParserRegistry().resolve_media_type(
            filename="old-format.doc",
            declared_media_type="application/msword",
        )
