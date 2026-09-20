from __future__ import annotations

import os
import re
import subprocess
import tempfile
from io import BytesIO
from pathlib import Path
from typing import Any

from docx import Document
from docx.shared import Cm, Pt
from fastapi import FastAPI, HTTPException, Response
from markdown_it import MarkdownIt
from pydantic import BaseModel, Field


RENDERER_VERSION = "document-renderer-v2"
DOCX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
PDF_MEDIA_TYPE = "application/pdf"
SOFFICE_BINARY = os.getenv("SOFFICE_BINARY", "soffice")
PDF_TIMEOUT_SECONDS = int(os.getenv("PDF_RENDER_TIMEOUT_SECONDS", "120"))
TEMPLATE_DIR = Path(os.getenv("DOCUMENT_RENDERER_TEMPLATE_DIR", "/opt/document-renderer/templates"))
SAFE_TEMPLATE = re.compile(r"^[A-Za-z0-9][A-Za-z0-9._-]{0,127}$")


class RenderDocxRequest(BaseModel):
    markdown: str = Field(min_length=1)
    title: str = Field(default="Propuesta", min_length=1)
    language: str = "es"
    template_id: str = "builtin-neutral"
    metadata: dict[str, Any] = Field(default_factory=dict)


app = FastAPI(title="Proposal Document Renderer", version="0.1.0")


@app.get("/health/live")
@app.get("/health/ready")
def health() -> dict[str, str]:
    return {"status": "UP", "renderer": RENDERER_VERSION}


@app.post("/v1/render/docx")
def render_docx(request: RenderDocxRequest) -> Response:
    try:
        payload = _render_docx_bytes(request)
        file_name = _safe_filename(request.title) + ".docx"
        return Response(content=payload, media_type=DOCX_MEDIA_TYPE, headers=_headers(file_name, request.template_id))
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(status_code=500, detail=f"DOCX rendering failed: {exc}") from exc


@app.post("/v1/render/pdf")
def render_pdf(request: RenderDocxRequest) -> Response:
    try:
        docx = _render_docx_bytes(request)
        pdf = _convert_docx_to_pdf(docx)
        file_name = _safe_filename(request.title) + ".pdf"
        return Response(content=pdf, media_type=PDF_MEDIA_TYPE, headers=_headers(file_name, request.template_id))
    except HTTPException:
        raise
    except subprocess.TimeoutExpired as exc:
        raise HTTPException(status_code=504, detail=f"PDF rendering timed out after {PDF_TIMEOUT_SECONDS}s") from exc
    except Exception as exc:
        raise HTTPException(status_code=500, detail=f"PDF rendering failed: {exc}") from exc


def _headers(file_name: str, template_id: str) -> dict[str, str]:
    return {
        "Content-Disposition": f'attachment; filename="{file_name}"',
        "X-Renderer-Version": RENDERER_VERSION,
        "X-Template-Id": template_id or "none",
    }


def _render_docx_bytes(request: RenderDocxRequest) -> bytes:
    document = _load_template(request.template_id)
    _prepare_document(document, request)
    _render_markdown(document, request.markdown)
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _convert_docx_to_pdf(payload: bytes) -> bytes:
    with tempfile.TemporaryDirectory(prefix="proposal-render-") as job_dir_raw:
        job_dir = Path(job_dir_raw)
        input_docx = job_dir / "document.docx"
        output_dir = job_dir / "output"
        profile_dir = job_dir / "lo-profile"
        output_dir.mkdir()
        profile_dir.mkdir()
        input_docx.write_bytes(payload)

        completed = subprocess.run(
            [
                SOFFICE_BINARY,
                f"-env:UserInstallation={profile_dir.as_uri()}",
                "--headless",
                "--nologo",
                "--nodefault",
                "--nofirststartwizard",
                "--convert-to", "pdf",
                "--outdir", str(output_dir),
                str(input_docx),
            ],
            stdout=subprocess.PIPE,
            stderr=subprocess.STDOUT,
            timeout=PDF_TIMEOUT_SECONDS,
            check=False,
            text=True,
        )
        pdf_path = output_dir / "document.pdf"
        if completed.returncode != 0 or not pdf_path.exists():
            detail = completed.stdout.strip()[-1500:] if completed.stdout else "no LibreOffice output"
            raise RuntimeError(f"LibreOffice conversion failed ({completed.returncode}): {detail}")
        return pdf_path.read_bytes()


def _load_template(template_id: str) -> Document:
    value = (template_id or "").strip()
    if value in {"", "builtin-neutral"}:
        return Document()
    if not SAFE_TEMPLATE.fullmatch(value):
        raise HTTPException(status_code=422, detail="Invalid proposal template identifier")
    path = TEMPLATE_DIR / f"{value}.docx"
    if not path.exists():
        raise HTTPException(status_code=404, detail=f"Proposal template '{value}' not found")
    document = Document(path)
    _clear_body_keep_section(document)
    return document


def _clear_body_keep_section(document: Document) -> None:
    body = document._element.body
    for element in list(body):
        if element.tag.endswith("}sectPr"):
            continue
        body.remove(element)


def _prepare_document(document: Document, request: RenderDocxRequest) -> None:
    props = document.core_properties
    props.title = request.title
    props.subject = "ProposalFlow generated proposal"
    props.comments = f"Rendered by {RENDERER_VERSION}; template={request.template_id or 'builtin-neutral'}"

    section = document.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.3)
    section.right_margin = Cm(2.3)

    normal = document.styles["Normal"]
    normal.font.name = "Liberation Sans"
    normal.font.size = Pt(10.5)

    for style_name in ("Title", "Heading 1", "Heading 2", "Heading 3"):
        if style_name in document.styles:
            document.styles[style_name].font.name = "Liberation Sans"

    if "Title" in document.styles:
        document.styles["Title"].font.size = Pt(24)


def _render_markdown(document: Document, markdown: str) -> None:
    parser = MarkdownIt("commonmark").enable("table")
    tokens = parser.parse(markdown)

    list_stack: list[str] = []
    i = 0
    while i < len(tokens):
        token = tokens[i]

        if token.type == "heading_open":
            level = int(token.tag[1:])
            inline = tokens[i + 1] if i + 1 < len(tokens) else None
            if level == 1:
                paragraph = document.add_paragraph(style="Title")
            else:
                style = f"Heading {min(level - 1, 3)}"
                paragraph = document.add_paragraph(style=style)
            _render_inline(paragraph, inline)
            i += 3
            continue

        if token.type == "bullet_list_open":
            list_stack.append("bullet")
            i += 1
            continue
        if token.type == "ordered_list_open":
            list_stack.append("number")
            i += 1
            continue
        if token.type in {"bullet_list_close", "ordered_list_close"}:
            if list_stack:
                list_stack.pop()
            i += 1
            continue

        if token.type == "paragraph_open":
            inline = tokens[i + 1] if i + 1 < len(tokens) else None
            style = None
            if list_stack:
                style = "List Bullet" if list_stack[-1] == "bullet" else "List Number"
            paragraph = document.add_paragraph(style=style)
            _render_inline(paragraph, inline)
            i += 3
            continue

        if token.type == "blockquote_open":
            i += 1
            while i < len(tokens) and tokens[i].type != "blockquote_close":
                if tokens[i].type == "paragraph_open":
                    inline = tokens[i + 1] if i + 1 < len(tokens) else None
                    paragraph = document.add_paragraph()
                    paragraph.paragraph_format.left_indent = Cm(0.8)
                    paragraph.paragraph_format.right_indent = Cm(0.4)
                    _render_inline(paragraph, inline)
                    i += 3
                else:
                    i += 1
            i += 1
            continue

        if token.type in {"fence", "code_block"}:
            paragraph = document.add_paragraph()
            run = paragraph.add_run(token.content.rstrip())
            run.font.name = "Liberation Mono"
            run.font.size = Pt(9)
            i += 1
            continue

        if token.type == "table_open":
            i = _render_table(document, tokens, i)
            continue

        if token.type == "hr":
            paragraph = document.add_paragraph()
            paragraph.add_run("─" * 72)
            i += 1
            continue

        i += 1


def _render_table(document: Document, tokens, start: int) -> int:
    rows: list[list[str]] = []
    current_row: list[str] | None = None
    i = start + 1
    while i < len(tokens) and tokens[i].type != "table_close":
        token = tokens[i]
        if token.type == "tr_open":
            current_row = []
        elif token.type in {"th_open", "td_open"}:
            inline = tokens[i + 1] if i + 1 < len(tokens) and tokens[i + 1].type == "inline" else None
            if current_row is not None:
                current_row.append(_inline_text(inline))
        elif token.type == "tr_close" and current_row is not None:
            rows.append(current_row)
            current_row = None
        i += 1

    if rows:
        width = max(len(row) for row in rows)
        table = document.add_table(rows=len(rows), cols=width)
        table.style = "Table Grid"
        for row_index, values in enumerate(rows):
            for col_index in range(width):
                cell = table.cell(row_index, col_index)
                cell.text = values[col_index] if col_index < len(values) else ""
                if row_index == 0:
                    for run in cell.paragraphs[0].runs:
                        run.bold = True
    return i + 1


def _render_inline(paragraph, inline) -> None:
    if inline is None:
        return
    strong = False
    emphasis = False
    link = False
    for child in inline.children or []:
        if child.type == "strong_open":
            strong = True
        elif child.type == "strong_close":
            strong = False
        elif child.type == "em_open":
            emphasis = True
        elif child.type == "em_close":
            emphasis = False
        elif child.type == "link_open":
            link = True
        elif child.type == "link_close":
            link = False
        elif child.type in {"softbreak", "hardbreak"}:
            paragraph.add_run().add_break()
        elif child.type == "code_inline":
            run = paragraph.add_run(child.content)
            run.font.name = "Liberation Mono"
            run.font.size = Pt(9)
        elif child.type in {"text", "html_inline"}:
            run = paragraph.add_run(child.content)
            run.bold = strong
            run.italic = emphasis
            if link:
                run.underline = True


def _inline_text(inline) -> str:
    if inline is None:
        return ""
    pieces: list[str] = []
    for child in inline.children or []:
        if child.type in {"text", "code_inline", "html_inline"}:
            pieces.append(child.content)
        elif child.type in {"softbreak", "hardbreak"}:
            pieces.append("\n")
    return "".join(pieces)


def _safe_filename(title: str) -> str:
    value = re.sub(r"[^A-Za-z0-9À-ÖØ-öø-ÿ._ -]+", "", title).strip()
    value = re.sub(r"\s+", "-", value)
    return (value or "proposal")[:120]
